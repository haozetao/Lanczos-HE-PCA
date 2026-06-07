from pathlib import Path

import fitz
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path("/Users/bytedance/PCA")
WORK = ROOT / "workspace" / "nudge-ppt"
OUT = WORK / "outputs-editable"
MEDIA = WORK / "template-media" / "2026-5-9"
PDF = ROOT / "paper" / "Henzinger 等 - Nudge A Private Recommendations Engine.pdf"

BLUE = RGBColor(0x00, 0x20, 0x60)
RED = RGBColor(0xC0, 0x00, 0x00)
BLACK = RGBColor(0x1C, 0x20, 0x28)
MUTED = RGBColor(0x5B, 0x63, 0x72)
LIGHT = RGBColor(0xF4, 0xF6, 0xFA)
LIGHT2 = RGBColor(0xEA, 0xEF, 0xF7)
LINE = RGBColor(0xD2, 0xD8, 0xE2)
DARK = RGBColor(0x16, 0x1C, 0x2A)
GREEN = RGBColor(0x2E, 0x7D, 0x50)
GOLD = RGBColor(0xD2, 0x97, 0x28)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT_CN = "微软雅黑"
FONT_EN = "Arial"


def rgb_tuple(c):
    return (c[0], c[1], c[2])


def set_run(run, size=18, bold=False, color=BLACK, font=FONT_CN):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def set_text(shape, text, size=18, bold=False, color=BLACK, align=PP_ALIGN.LEFT, font=FONT_CN):
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return shape


def add_text(slide, x, y, w, h, text, size=18, bold=False, color=BLACK, align=PP_ALIGN.LEFT, font=FONT_CN):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.margin_left = Inches(0.02)
    box.text_frame.margin_right = Inches(0.02)
    box.text_frame.margin_top = Inches(0.02)
    box.text_frame.margin_bottom = Inches(0.02)
    return set_text(box, text, size=size, bold=bold, color=color, align=align, font=font)


def add_multiline(slide, x, y, w, h, paragraphs, size=16, color=BLACK, bullet=False, gap=4):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = FONT_CN
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
        if bullet:
            p.level = 0
            p._p.get_or_add_pPr().set("marL", "285750")
            p._p.get_or_add_pPr().set("indent", "-171450")
    return box


def add_card(slide, x, y, w, h, title, body=None, accent=BLUE, fill=WHITE):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = LINE
    shp.line.width = Pt(1)
    add_text(slide, x + 0.18, y + 0.16, w - 0.36, 0.32, title, size=15, bold=True, color=accent)
    if body:
        add_multiline(slide, x + 0.18, y + 0.58, w - 0.36, h - 0.72, body if isinstance(body, list) else [body], size=12, color=BLACK, bullet=isinstance(body, list))
    return shp


def add_pill(slide, x, y, text, color=BLUE, w=None):
    w = w or max(1.1, len(text) * 0.12 + 0.35)
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.32))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.color.rgb = color
    set_text(shp, text, size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    return shp


def add_arrow(slide, x1, y1, x2, y2, color=BLUE, width=2):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    conn.line.end_arrowhead = True
    return conn


def crop_pdf(page_num, rel, name):
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / name
    if path.exists():
        return path
    doc = fitz.open(str(PDF))
    page = doc[page_num - 1]
    rect = page.rect
    clip = fitz.Rect(
        rect.x0 + rel[0] * rect.width,
        rect.y0 + rel[1] * rect.height,
        rect.x0 + rel[2] * rect.width,
        rect.y0 + rel[3] * rect.height,
    )
    pix = page.get_pixmap(matrix=fitz.Matrix(3, 3), clip=clip, alpha=False)
    pix.save(str(path))
    return path


def add_header(slide, title, section=None, page=None):
    # Top reference-template bar.
    slide.shapes.add_picture(str(MEDIA / "image1.png"), Inches(0), Inches(0), width=Inches(13.33), height=Inches(0.68))
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(0.12), Inches(0.68))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.color.rgb = BLUE
    # Right logo as image + editable labels.
    slide.shapes.add_picture(str(MEDIA / "image7.jpeg"), Inches(10.72), Inches(0.08), width=Inches(0.42), height=Inches(0.42))
    add_text(slide, 11.18, 0.08, 1.2, 0.18, "华中科技大学", size=8, bold=True, color=BLACK)
    add_text(slide, 11.18, 0.28, 1.65, 0.18, "网络空间安全学院", size=7, color=MUTED)
    if section:
        add_text(slide, 0.70, 0.92, 2.5, 0.24, section, size=8, bold=True, color=RED)
    add_text(slide, 0.70, 1.13, 11.7, 0.5, title, size=23, bold=True, color=BLUE)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.70), Inches(1.68), Inches(11.95), Inches(0.015))
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.color.rgb = LINE
    add_text(slide, 0.70, 7.08, 4.0, 0.2, "Nudge: A Private Recommendations Engine", size=7, color=MUTED)
    if page:
        add_text(slide, 12.45, 7.08, 0.35, 0.2, f"{page:02d}", size=7, color=MUTED, align=PP_ALIGN.RIGHT)


def blank(prs, title, section=None, page=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, title, section, page)
    return slide


def cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(MEDIA / "image6.jpeg"), Inches(0), Inches(0), width=Inches(13.33), height=Inches(7.5))
    wash = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(7.5))
    wash.fill.solid()
    wash.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    wash.fill.transparency = 25
    wash.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    add_header(slide, "", None, None)
    add_text(slide, 0.85, 1.85, 4.2, 0.8, "Nudge", size=44, bold=True, color=RED, font=FONT_EN)
    add_text(slide, 0.88, 2.55, 8.8, 0.55, "A Private Recommendations Engine", size=27, bold=True, color=BLUE, font=FONT_EN)
    add_text(slide, 0.90, 3.22, 9.0, 0.35, "Henzinger · Dauterman · Corrigan-Gibbs · Boneh", size=16, color=BLACK, font=FONT_EN)
    add_text(slide, 0.90, 3.70, 7.8, 0.34, "论文汇报：三方 MPC 下的大规模私有推荐系统", size=17, bold=True, color=BLUE)
    add_pill(slide, 0.90, 6.10, "Recommendation", BLUE, 1.55)
    add_pill(slide, 2.60, 6.10, "3PC", RED, 0.75)
    add_pill(slide, 3.50, 6.10, "Function Secret Sharing", BLUE, 1.95)
    add_pill(slide, 5.60, 6.10, "Power Iteration", RED, 1.45)
    add_card(slide, 8.80, 2.0, 3.55, 2.85, "核心问题", [
        "如何在不泄露用户评分的情况下训练推荐模型？",
        "如何让密码学私有推荐扩展到 Netflix 量级？",
    ], accent=RED, fill=RGBColor(0xFF, 0xFF, 0xFF))
    return slide


def slide_motivation(prs):
    slide = blank(prs, "为什么需要 Nudge？", "一、研究背景", 2)
    add_card(slide, 0.82, 2.0, 3.55, 1.05, "推荐系统的隐私矛盾", "点击、浏览、评分等交互数据高度敏感，却正是个性化推荐的燃料。", RED)
    add_card(slide, 0.82, 3.35, 3.55, 1.05, "现有方案的不足", "匿名、联邦学习、TEE、传统 MPC 都在泄露面、信任假设或规模上受限。", BLUE)
    add_card(slide, 0.82, 4.70, 3.55, 1.05, "Nudge 的目标", "在强密码学隐私下，仍能训练有用的推荐模型并支撑大规模数据。", GREEN)
    dark = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(5.1), Inches(2.0), Inches(6.65), Inches(3.75))
    dark.fill.solid()
    dark.fill.fore_color.rgb = DARK
    dark.line.color.rgb = DARK
    add_text(slide, 5.45, 2.32, 3.2, 0.35, "论文给出的关键结果", size=18, bold=True, color=WHITE)
    add_multiline(slide, 5.50, 2.92, 5.75, 2.25, [
        "Netflix：463,435 用户、17,769 电影、56M ratings",
        "三台 192-core 服务器，LAN 下 50 min 完成私有训练",
        "server-to-server 通信约 40 GB",
        "nDCG@20 = 0.29，接近明文矩阵分解，与神经方法 0.31 相近",
    ], size=14, color=WHITE, bullet=True)
    add_text(slide, 5.45, 5.25, 5.7, 0.3, "一句话：不是追求最复杂模型，而是让“可扩展的私有推荐”第一次变得现实。", size=13, bold=True, color=RGBColor(0xF2, 0xB1, 0x92))


def slide_architecture(prs):
    slide = blank(prs, "系统总览：三服务器拆分信任", "二、系统设计", 3)
    fig1 = crop_pdf(4, (0.04, 0.08, 0.47, 0.34), "paper-fig1-protocol-flow.png")
    slide.shapes.add_picture(str(fig1), Inches(0.85), Inches(2.0), width=Inches(4.4))
    add_text(slide, 0.85, 4.95, 4.4, 0.3, "论文 Figure 1：Nudge 协议流", size=10, color=MUTED, align=PP_ALIGN.CENTER)
    steps = [
        ("1", "Private data collection", "用户把评分向量 secret-share 后分别发给三台服务器。"),
        ("2", "Private matrix factorization", "服务器在 3PC 中训练 user embeddings A 与 item embeddings B。"),
        ("3", "Private recommendations", "服务器返回推荐分数的 share，用户本地重构推荐列表。"),
    ]
    y = 2.0
    for no, title, desc in steps:
        circ = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(6.0), Inches(y), Inches(0.42), Inches(0.42))
        circ.fill.solid()
        circ.fill.fore_color.rgb = BLUE
        circ.line.color.rgb = BLUE
        set_text(circ, no, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_text(slide, 6.55, y - 0.02, 4.5, 0.28, title, size=15, bold=True, color=BLACK, font=FONT_EN)
        add_text(slide, 6.55, y + 0.38, 5.3, 0.52, desc, size=12, color=MUTED)
        y += 1.25
    add_card(slide, 6.0, 5.75, 5.65, 0.75, "与联邦学习不同", "用户不参与多轮训练；训练过程也不泄露中间模型版本。", RED, fill=RGBColor(0xFB, 0xF3, 0xEF))


def slide_security(prs):
    slide = blank(prs, "威胁模型：强隐私，但不是万能", "二、系统设计", 4)
    cx, cy = 3.35, 3.85
    coords = [(3.35, 2.15, "S0"), (1.95, 4.75, "S1"), (4.75, 4.75, "S2")]
    for x, y, label in coords:
        add_arrow(slide, cx, cy, x + 0.25, y + 0.25, LINE, 1.5)
        node = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(0.62), Inches(0.62))
        node.fill.solid()
        node.fill.fore_color.rgb = DARK
        node.line.color.rgb = DARK
        set_text(node, label, size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    center = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(cx - 0.55), Inches(cy - 0.55), Inches(1.1), Inches(1.1))
    center.fill.solid()
    center.fill.fore_color.rgb = BLUE
    center.line.color.rgb = BLUE
    set_text(center, "2-out-of-3\nsharing", size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    cards = [
        ("保护对象", "隐藏用户评分矩阵 U 的位置和值；其他用户也看不到单个用户行为。", GREEN),
        ("允许泄露", "最终聚合模型 B 所编码的信息；每个用户提交了多少个非零评分。", GOLD),
        ("不保护", "两台及以上服务器串谋；恶意服务器偏离协议；时间等元数据。", RED),
    ]
    y = 2.0
    for title, body, color in cards:
        add_card(slide, 7.0, y, 4.9, 0.95, title, body, color)
        y += 1.25
    add_text(slide, 0.95, 6.15, 10.8, 0.35, "安全目标：against one honest-but-curious server colluding with malicious users.", size=14, bold=True, color=RED, font=FONT_EN)


def slide_mf(prs):
    slide = blank(prs, "矩阵分解直觉：把评分矩阵压成低秩表示", "三、算法直觉", 5)
    add_text(slide, 0.95, 2.0, 2.4, 0.28, "用户 × 物品评分矩阵 U", size=14, bold=True, color=BLUE)
    # Matrix U with native rectangles.
    for i in range(8):
        for j in range(10):
            val = (i * 7 + j * 11) % 10
            color = RGBColor(0xF0, 0xF4, 0xFA) if val < 5 else RGBColor(0x6B, 0xA2, 0xC8)
            shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.95 + j * 0.26), Inches(2.50 + i * 0.22), Inches(0.21), Inches(0.16))
            shp.fill.solid()
            shp.fill.fore_color.rgb = color
            shp.line.color.rgb = WHITE
    add_text(slide, 4.0, 3.05, 0.35, 0.4, "≈", size=32, bold=True, color=RED)
    add_card(slide, 4.65, 2.45, 1.8, 2.2, "A", ["用户嵌入", "m × d"], BLUE, fill=RGBColor(0xF1, 0xF8, 0xF4))
    add_text(slide, 6.85, 3.05, 0.35, 0.4, "×", size=28, bold=True, color=RED)
    add_card(slide, 7.55, 2.60, 3.35, 1.75, "B", ["物品嵌入 / 推荐模型", "d × n，行正交"], RED, fill=RGBColor(0xFB, 0xF3, 0xEF))
    add_card(slide, 0.95, 5.45, 10.4, 0.85, "服务阶段怎么推荐？", "对用户 i，推荐分数等价于 u(i) · BᵀB；选择分数最高且用户未看过的物品。", BLUE, fill=RGBColor(0xF7, 0xF9, 0xFC))
    add_text(slide, 0.95, 6.48, 10.6, 0.32, "Nudge 的隐私边界也在这里：服务器最终学习到的是聚合模型 B，而不是单个用户的评分向量。", size=13, bold=True, color=RED)


def slide_algorithm(prs):
    slide = blank(prs, "算法步骤：Power Iteration 做私有矩阵分解", "四、核心算法", 6)
    fig4 = crop_pdf(7, (0.05, 0.05, 0.56, 0.37), "paper-fig4-power-iteration.png")
    slide.shapes.add_picture(str(fig4), Inches(0.85), Inches(2.0), width=Inches(4.65))
    add_text(slide, 0.85, 5.10, 4.65, 0.25, "论文 Figure 4：ApproxFactor(U)", size=10, color=MUTED, align=PP_ALIGN.CENTER)
    steps = [
        ("初始化", "B 置零；对第 i 个 latent factor 随机初始化向量 v。"),
        ("正交化", "SetOrthogonal(v, B)：从 v 中移除已求方向分量。"),
        ("幂迭代", "重复 ℓ 次：v ← Uᵀ(Uv)，然后正交化并归一化。"),
        ("写入模型", "将收敛后的 v 作为 B 的第 i 行；继续提取下一个因子。"),
        ("线性恢复", "最后计算 A := U·Bᵀ；推荐分数来自 A·B。"),
    ]
    y = 1.95
    for idx, (title, body) in enumerate(steps, 1):
        color = BLUE if idx <= 3 else RED
        circ = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(6.1), Inches(y), Inches(0.38), Inches(0.38))
        circ.fill.solid()
        circ.fill.fore_color.rgb = color
        circ.line.color.rgb = color
        set_text(circ, str(idx), size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_text(slide, 6.65, y - 0.02, 4.6, 0.25, title, size=14, bold=True, color=BLACK)
        add_text(slide, 6.65, y + 0.30, 5.0, 0.36, body, size=11.5, color=MUTED)
        if idx < len(steps):
            add_arrow(slide, 6.29, y + 0.44, 6.29, y + 0.76, LINE, 1)
        y += 0.88


def slide_mpc(prs):
    slide = blank(prs, "为什么这个算法适合 3PC？", "四、核心算法", 7)
    add_card(slide, 0.85, 2.0, 3.0, 3.0, "线性大块", [
        "矩阵-向量乘法：U·v、Uᵀ·v",
        "Replicated secret sharing 可非交互高速完成",
        "大矩阵计算主成本接近明文 3×",
    ], BLUE)
    add_arrow(slide, 3.95, 3.5, 4.75, 3.5, BLUE, 2)
    add_card(slide, 4.95, 2.0, 3.0, 3.0, "非线性小块", [
        "Truncₜ：固定点乘法后的截断",
        "Normalize：计算 1/||v||",
        "这些是通信和轮数的关键瓶颈",
    ], RED)
    add_arrow(slide, 8.05, 3.5, 8.85, 3.5, BLUE, 2)
    dark = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(9.05), Inches(2.0), Inches(3.2), Inches(3.0))
    dark.fill.solid()
    dark.fill.fore_color.rgb = DARK
    dark.line.color.rgb = DARK
    add_text(slide, 9.28, 2.25, 2.5, 0.28, "Nudge 优化", size=16, bold=True, color=WHITE)
    add_multiline(slide, 9.28, 2.75, 2.55, 1.75, [
        "FSS 处理比较 / 最高有效位",
        "截断通信降低约 6×",
        "归一化用 Newton-Raphson",
    ], size=12, color=WHITE, bullet=True)
    add_card(slide, 1.1, 5.75, 10.6, 0.75, "核心 co-design", "把算法写成“矩阵-向量程序”：大规模线性代数尽量不通信，只在必要的非线性函数上支付 3PC 成本。", RED, fill=RGBColor(0xFB, 0xF3, 0xEF))


def slide_complexity(prs):
    slide = blank(prs, "复杂度叙事：通信从“跟评分数走”变成“跟用户+物品走”", "五、性能分析", 8)
    add_text(slide, 0.95, 2.0, 4.8, 0.32, "固定 embedding 维度 d 时：", size=17, bold=True, color=BLACK)
    add_card(slide, 0.95, 2.65, 4.4, 1.25, "Compute", "O(mn)：密态/secret-shared 矩阵看不到稀疏结构，必须按密集矩阵处理。", BLUE)
    add_card(slide, 0.95, 4.35, 4.4, 1.25, "Communication", "O(m+n)：通信不随评分数线性增长，这是相比许多 prior systems 的关键优势。", RED)
    add_text(slide, 7.0, 2.1, 2.8, 0.3, "通信量级直觉", size=17, bold=True, color=BLACK)
    # Native bar chart.
    for y, label, frac, color in [(3.0, "传统私有 MF", 0.92, RED), (4.25, "Nudge", 0.18, BLUE)]:
        add_text(slide, 6.65, y - 0.32, 2.0, 0.22, label, size=13, bold=True, color=BLACK)
        bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(6.65), Inches(y), Inches(4.2), Inches(0.30))
        bg.fill.solid()
        bg.fill.fore_color.rgb = LIGHT2
        bg.line.color.rgb = LIGHT2
        bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(6.65), Inches(y), Inches(4.2 * frac), Inches(0.30))
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.color.rgb = color
    add_card(slide, 6.65, 5.55, 4.75, 0.82, "瓶颈仍在计算", "当用户和物品都极大时，Nudge 需要随机投影等降维扩展来突破 O(mn)。", RED, fill=RGBColor(0xFB, 0xF3, 0xEF))


def slide_quality(prs):
    slide = blank(prs, "推荐质量：接近明文矩阵分解", "六、实验结果", 9)
    fig5 = crop_pdf(9, (0.05, 0.05, 0.62, 0.36), "paper-fig5-quality.png")
    slide.shapes.add_picture(str(fig5), Inches(0.85), Inches(2.0), width=Inches(5.45))
    add_text(slide, 0.85, 5.35, 5.45, 0.25, "论文 Figure 5：Netflix 上 Recall / nDCG 对比", size=10, color=MUTED, align=PP_ALIGN.CENTER)
    metrics = [("Nudge nDCG@20", "0.29", BLUE), ("MultVAE nDCG@20", "0.31", RED), ("Nudge Recall@20", "0.23", GREEN)]
    y = 2.12
    for label, val, color in metrics:
        card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(y), Inches(4.1), Inches(0.68))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = LINE
        add_text(slide, 7.25, y + 0.20, 2.3, 0.22, label, size=12, bold=True, color=MUTED, font=FONT_EN)
        add_text(slide, 10.15, y + 0.12, 0.7, 0.34, val, size=20, bold=True, color=color, align=PP_ALIGN.RIGHT, font=FONT_EN)
        y += 0.95
    add_card(slide, 7.0, 5.05, 4.1, 0.95, "实验结论", "固定步数幂迭代、固定点算术、延迟截断与近似归一化，并没有显著伤害推荐质量。", RED, fill=RGBColor(0xFB, 0xF3, 0xEF))


def slide_performance(prs):
    slide = blank(prs, "性能：私有矩阵分解从“天/小时”压到“秒/分钟”", "六、实验结果", 10)
    table6 = crop_pdf(10, (0.03, 0.05, 0.97, 0.43), "paper-table6-performance.png")
    slide.shapes.add_picture(str(table6), Inches(0.75), Inches(1.95), width=Inches(11.7))
    add_text(slide, 0.75, 4.55, 11.7, 0.2, "论文 Table 6：与 prior private matrix factorization / MP-SPDZ baseline 对比", size=9, color=MUTED, align=PP_ALIGN.CENTER)
    highlights = [
        ("MovieLens100K", "10.7 sec LAN / 51.8 sec WAN", "0.5 GB"),
        ("MovieLens1M", "32.8 sec LAN / 1.8 min WAN", "1.4 GB"),
        ("Netflix", "50.2 min LAN / 1h13 WAN", "38.9 GB"),
    ]
    x = 0.95
    for name, time, comm in highlights:
        add_card(slide, x, 5.18, 3.55, 0.92, name, [time, f"server-to-server: {comm}"], BLUE)
        x += 4.05
    add_text(slide, 0.95, 6.48, 10.8, 0.26, "论文强调：相比 garbled-circuit 系统，通信和计算低四个数量级；相比 PIRSONA，通信少 8× 且少 1 个非串谋方。", size=12, bold=True, color=RED)


def slide_scale(prs):
    slide = blank(prs, "端到端与横向扩展：训练重，服务轻", "六、实验结果", 11)
    fig7 = crop_pdf(11, (0.05, 0.05, 0.96, 0.34), "paper-fig7-scale.png")
    table8 = crop_pdf(12, (0.05, 0.05, 0.56, 0.33), "paper-table8-serving.png")
    slide.shapes.add_picture(str(fig7), Inches(0.75), Inches(1.95), width=Inches(6.9))
    slide.shapes.add_picture(str(table8), Inches(8.0), Inches(1.95), width=Inches(3.9))
    add_text(slide, 0.75, 4.60, 6.9, 0.2, "Figure 7：百万级用户/物品下的扩展趋势", size=9, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, 8.0, 4.60, 3.9, 0.2, "Table 8：Netflix 上端到端服务开销", size=9, color=MUTED, align=PP_ALIGN.CENTER)
    cards = [
        ("Criteo", "6.1M users / 700 items", "1.5 h · 125 GB"),
        ("Yelp", "279K users / 148K items", "8 h · 248 GB"),
        ("Serving", "Netflix user fetch", "298 KB · 0.38 s"),
    ]
    x = 0.95
    for title, mid, bottom in cards:
        add_card(slide, x, 5.30, 3.55, 0.90, title, [mid, bottom], RED if title == "Serving" else BLUE, fill=RGBColor(0xF7, 0xF9, 0xFC))
        x += 4.05
    add_text(slide, 0.95, 6.48, 10.9, 0.25, "读法：训练阶段是重计算；但模型训练好后，用户拉取推荐只需数百 KB、亚秒级延迟。", size=12, bold=True, color=BLUE)


def slide_takeaways(prs):
    slide = blank(prs, "局限、扩展与汇报结论", "七、总结", 12)
    cols = [
        ("局限", [
            "需要 3 个非串谋服务器，只容忍 1 个半诚实 compromise",
            "不负责私有内容获取；需和 PIR / 匿名代理配合",
            "最终模型 B 仍可能编码聚合用户行为信息",
            "训练成本显著高于明文系统，且密态矩阵无法利用稀疏性",
        ], RED),
        ("扩展", [
            "随机投影：先把 U·S 压到低维 sketch，再做私有矩阵分解",
            "差分隐私：对 UᵀU 加高斯噪声，Netflix nDCG@20 从 0.29 降到 0.24",
            "输入有效性检查：验证评分在 {0,1}，几乎免费",
            "可迁移到 PCA、PageRank、谱聚类、低维 embedding 等任务",
        ], BLUE),
        ("一句话评价", [
            "Nudge 的核心贡献不是“更复杂的推荐模型”，而是选择了适合 MPC 的线性代数算法",
            "通过 power iteration + 3PC/FSS co-design，把私有推荐推进到 Netflix 量级",
            "适合从“算法结构如何决定密码协议成本”这个角度讲",
        ], GREEN),
    ]
    x = 0.75
    for title, items, color in cols:
        add_card(slide, x, 2.0, 3.75, 4.25, title, items, color)
        x += 4.12
    add_text(slide, 0.95, 6.55, 10.8, 0.28, "开放问题：恶意安全、多方数扩展、仅 2 个非串谋服务器且不牺牲性能。", size=13, bold=True, color=RED)


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    cover(prs)
    slide_motivation(prs)
    slide_architecture(prs)
    slide_security(prs)
    slide_mf(prs)
    slide_algorithm(prs)
    slide_mpc(prs)
    slide_complexity(prs)
    slide_quality(prs)
    slide_performance(prs)
    slide_scale(prs)
    slide_takeaways(prs)
    out = OUT / "nudge-private-recommendations-engine-editable-template.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build()
